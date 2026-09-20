# planner.py
import os
from structure_extractor import extract_structure


def _read_workbook_state(workbook_path):
    """Read current sheets + columns from the workbook."""
    if not workbook_path or not os.path.exists(workbook_path):
        return {"sheets": {}, "columns": {}}

    try:
        from openpyxl import load_workbook
        wb = load_workbook(workbook_path, data_only=False)

        sheets = {}
        for ws in wb.worksheets:
            # Read headers (row 1)
            headers = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=1, column=c).value
                if v is not None:
                    headers.append(str(v).strip())

            # Count rows
            row_count = 0
            for r in range(2, ws.max_row + 1):
                row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                if any(v not in (None, "") for v in row_vals):
                    row_count += 1

            sheets[ws.title] = {
                "headers": headers,
                "rows": row_count,
            }

        return {"sheets": sheets}
    except Exception as e:
        print("[planner] read state error:", e)
        return {"sheets": {}}


def build_plan_for_task(task, workbook_path, session=None):
    task = (task or "").strip()
    if not task:
        return {"success": False, "message": "No task provided."}

    state = _read_workbook_state(workbook_path)
    existing_sheets = state.get("sheets", {})

    # ---- If session has PDF analysis → build from analysis ----
    if session:
        analysis = session.get("last_document_analysis") or {}
        if analysis.get("success"):
            cols = analysis.get("columns") or []
            rows = analysis.get("rows") or []
            sheet = analysis.get("sheet_name") or "ExtractedData"

            # If user said "Sheet2", use that
            import re
            m = re.search(r"\b(Sheet\s*\d+)\b", task, re.IGNORECASE)
            if m:
                sheet = m.group(1).strip()

            if cols and rows:
                # ✅ Check if the sheet already exists with these columns
                sheet_exists = sheet in existing_sheets
                existing_headers = existing_sheets.get(sheet, {}).get("headers", [])

                missing_cols = [
                    c for c in cols
                    if c.strip().lower() not in [h.lower() for h in existing_headers]
                ]

                steps = []

                # Only create the sheet if it doesn't exist
                if not sheet_exists:
                    steps.append({
                        "tool": "create_sheet",
                        "arguments": {"file_path": workbook_path, "sheet_name": sheet},
                    })

                # Only add missing columns
                for col in missing_cols:
                    steps.append({
                        "tool": "add_column",
                        "arguments": {"file_path": workbook_path,
                                      "sheet_name": sheet,
                                      "column_name": col},
                    })

                # Only add rows if the sheet has no data yet
                existing_rows = existing_sheets.get(sheet, {}).get("rows", 0)
                if existing_rows == 0:
                    for row in rows:
                        steps.append({
                            "tool": "add_row",
                            "arguments": {"file_path": workbook_path,
                                          "sheet_name": sheet,
                                          "row_data": row},
                        })

                # If nothing to do → already satisfied
                if not steps:
                    return {
                        "success": True,
                        "already_satisfied": True,
                        "summary": f"Sheet '{sheet}' already has all columns and data.",
                        "plan": {"steps": []},
                    }

                summary = (
                    f"• Sheet: {sheet}\n"
                    + (f"• Create sheet\n" if not sheet_exists else "")
                    + (f"• Add {len(missing_cols)} missing column(s)\n" if missing_cols else "")
                    + (f"• Add {len(rows)} row(s)\n" if existing_rows == 0 else "")
                )

                return {
                    "success": True,
                    "action": "create_sheet_with_data",
                    "summary": summary.strip(),
                    "plan": {"steps": steps},
                    "structure": analysis,
                    "needs_permission": True,
                }

    # ---- Otherwise use LLM ----
    doc_context = None
    if session:
        analysis = session.get("last_document_analysis") or {}
        if analysis.get("success"):
            doc_context = (
                f"Document type: {analysis.get('document_type')}\n"
                f"Columns: {analysis.get('columns')}\n"
                f"Rows: {analysis.get('rows')}\n"
                f"Existing sheets: {list(existing_sheets.keys())}"
            )

    structure = extract_structure(task, document_context=doc_context)
    if not structure.get("success"):
        return {"success": False, "message": f"Could not understand: {structure.get('error')}"}

    action = structure.get("action")
    sheet_name = structure.get("sheet_name", "")
    columns = structure.get("columns", [])
    rows = structure.get("rows", [])

    steps = []

    if action == "create_sheet_with_data":
        if not sheet_name or not columns:
            return {"success": False, "message": "Missing sheet name or columns."}

        sheet_exists = sheet_name in existing_sheets
        existing_headers = existing_sheets.get(sheet_name, {}).get("headers", [])

        missing_cols = [
            c for c in columns
            if c.strip().lower() not in [h.lower() for h in existing_headers]
        ]

        if not sheet_exists:
            steps.append({"tool": "create_sheet",
                          "arguments": {"file_path": workbook_path, "sheet_name": sheet_name}})

        for col in missing_cols:
            steps.append({"tool": "add_column",
                          "arguments": {"file_path": workbook_path,
                                        "sheet_name": sheet_name,
                                        "column_name": col}})

        existing_rows = existing_sheets.get(sheet_name, {}).get("rows", 0)
        if existing_rows == 0:
            for row in rows:
                steps.append({"tool": "add_row",
                              "arguments": {"file_path": workbook_path,
                                            "sheet_name": sheet_name,
                                            "row_data": row}})

    elif action == "add_row":
        if not sheet_name:
            return {"success": False, "message": "Missing sheet name."}
        for row in rows:
            steps.append({"tool": "add_row",
                          "arguments": {"file_path": workbook_path,
                                        "sheet_name": sheet_name,
                                        "row_data": row}})

    elif action == "add_column":
        if not sheet_name or not columns:
            return {"success": False, "message": "Missing sheet or columns."}
        for col in columns:
            steps.append({"tool": "add_column",
                          "arguments": {"file_path": workbook_path,
                                        "sheet_name": sheet_name,
                                        "column_name": col}})

    elif action == "edit_cell":
        if not sheet_name or not structure.get("cell"):
            return {"success": False, "message": "Missing sheet or cell."}
        steps.append({"tool": "edit_cell",
                      "arguments": {"file_path": workbook_path,
                                    "sheet_name": sheet_name,
                                    "cell": structure.get("cell"),
                                    "new_value": structure.get("value", "")}})

    elif action == "add_formula":
        col_name = structure.get("column_name", "Total")
        cell = structure.get("cell", "D2")
        formula = structure.get("formula", "")
        if not sheet_name or not formula:
            return {"success": False, "message": "Missing sheet or formula."}
        steps.append({"tool": "add_column",
                      "arguments": {"file_path": workbook_path,
                                    "sheet_name": sheet_name,
                                    "column_name": col_name}})
        steps.append({"tool": "write_formula",
                      "arguments": {"file_path": workbook_path,
                                    "sheet_name": sheet_name,
                                    "cell": cell,
                                    "formula": formula}})

    elif action == "format_cells":
        if not sheet_name:
            return {"success": False, "message": "Missing sheet."}
        cell_range = structure.get("cell_range", "A1:C1")
        steps.append({"tool": "format_cells",
                      "arguments": {"file_path": workbook_path,
                                    "sheet_name": sheet_name,
                                    "cell_range": cell_range,
                                    "bold": bool(structure.get("bold", True))}})

    elif action == "sort_data":
        if not sheet_name:
            return {"success": False, "message": "Missing sheet."}
        steps.append({"tool": "sort_data",
                      "arguments": {"file_path": workbook_path,
                                    "sheet_name": sheet_name,
                                    "column_number": int(structure.get("sort_column", 1)),
                                    "descending": bool(structure.get("descending", False))}})

    else:
        return {"success": False, "message": "Orbit could not understand the request."}

    if not steps:
        return {
            "success": True,
            "already_satisfied": True,
            "summary": "Nothing new to do — the workbook already matches the request.",
            "plan": {"steps": []},
        }

    summary_lines = []
    if action == "create_sheet_with_data":
        summary_lines.append(f"• Sheet '{sheet_name}'")
        if sheet_name not in existing_sheets:
            summary_lines.append(f"• Create sheet")
        if rows:
            summary_lines.append(f"• Add {len(rows)} row(s)")
    elif action == "add_row":
        summary_lines.append(f"• Add {len(rows)} row(s) to '{sheet_name}'")
    elif action == "add_column":
        summary_lines.append(f"• Add column(s) {', '.join(columns)}")
    elif action == "edit_cell":
        summary_lines.append(f"• Set {sheet_name}!{structure.get('cell')}")
    elif action == "add_formula":
        summary_lines.append(f"• Add formula '{structure.get('formula')}'")
    elif action == "format_cells":
        summary_lines.append(f"• Format '{structure.get('cell_range')}'")
    elif action == "sort_data":
        summary_lines.append(f"• Sort '{sheet_name}'")

    return {
        "success": True,
        "action": action,
        "summary": "\n".join(summary_lines),
        "plan": {"steps": steps},
        "structure": structure,
        "needs_permission": True,
    }