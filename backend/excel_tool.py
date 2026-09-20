# excel_tool.py
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
import os


# ============================================================
# CLOUD SAVE — module-level session info
# ============================================================

_CURRENT_SESSION_ID = None
_CURRENT_FILENAME = None


def set_cloud_context(session_id: str, filename: str = None):
    """Called by main.py / agent.py so saves also go to Supabase."""
    global _CURRENT_SESSION_ID, _CURRENT_FILENAME
    _CURRENT_SESSION_ID = session_id
    _CURRENT_FILENAME = filename


def _cloud_save(workbook, file_path):
    """
    Save locally AND upload to Supabase Storage (best-effort).
    Never fails the local save if cloud upload errors.
    """
    # 1. Local save (always)
    workbook.save(file_path)

    # 2. Cloud upload (best-effort)
    if _CURRENT_SESSION_ID:
        try:
            from file_store import upload_file
            filename = _CURRENT_FILENAME or os.path.basename(file_path)
            remote_path = f"{_CURRENT_SESSION_ID}/{filename}"
            result = upload_file(file_path, remote_path)
            if result.get("success"):
                print(f"[excel_tool] ☁️  Uploaded to Supabase: {remote_path}")
            else:
                print(f"[excel_tool] ⚠️  Cloud upload failed: {result.get('error')}")
        except Exception as e:
            print(f"[excel_tool] ⚠️  Cloud upload exception: {e}")


# ============================================================
# WORKBOOK FUNCTIONS
# ============================================================

def create_excel(file_path, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)

    for row_number, row_data in enumerate(rows, start=2):
        for column, value in enumerate(row_data, start=1):
            sheet.cell(row=row_number, column=column, value=value)

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "file": file_path,
        "sheet": "Data",
        "rows": len(rows),
        "columns": len(headers)
    }


def read_excel(file_path):
    workbook = load_workbook(file_path)
    result = {}

    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
        result[sheet.title] = rows

    return result


def read_range(file_path, sheet_name, start_cell, end_cell):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    rows = []

    for row in sheet.iter_rows(
        min_row=sheet[start_cell].row,
        max_row=sheet[end_cell].row,
        min_col=sheet[start_cell].column,
        max_col=sheet[end_cell].column,
        values_only=True
    ):
        rows.append(list(row))

    return {
        "sheet": sheet_name,
        "range": f"{start_cell}:{end_cell}",
        "data": rows
    }


def edit_cell(file_path, sheet_name, cell, new_value):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    old_value = sheet[cell].value
    sheet[cell] = new_value

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "cell": cell,
        "old_value": old_value,
        "new_value": new_value
    }


def add_row(file_path, sheet_name, row_data):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    if sheet.max_row == 1 and sheet["A1"].value is None:
        next_row = 1
    else:
        next_row = sheet.max_row + 1

    for column, value in enumerate(row_data, start=1):
        sheet.cell(
            row=next_row,
            column=column,
            value=value
        )

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "row": next_row,
        "data": row_data
    }


def add_column(file_path, sheet_name, column_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    sheet_is_empty = (
        sheet.max_row == 1
        and sheet.max_column == 1
        and sheet["A1"].value is None
    )

    if sheet_is_empty:
        next_column = 1
    else:
        next_column = sheet.max_column + 1

    sheet.cell(
        row=1,
        column=next_column,
        value=column_name
    )

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "column": next_column,
        "name": column_name
    }


def create_sheet(file_path, sheet_name):
    workbook = load_workbook(file_path)

    if sheet_name in workbook.sheetnames:
        return {
            "success": False,
            "created": False,
            "message": "Sheet already exists"
        }

    workbook.create_sheet(sheet_name)
    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "created": True,
        "sheet": sheet_name
    }


def write_range(file_path, sheet_name, start_cell, data):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    start_row, start_column = coordinate_to_tuple(start_cell)

    for row_offset, row_data in enumerate(data):
        for column_offset, value in enumerate(row_data):
            sheet.cell(
                row=start_row + row_offset,
                column=start_column + column_offset,
                value=value
            )

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "start_cell": start_cell,
        "rows_written": len(data)
    }


def write_formula(file_path, sheet_name, cell, formula):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    sheet[cell] = formula

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "cell": cell,
        "formula": formula
    }


def format_cells(
    file_path,
    sheet_name,
    cell_range,
    bold=False,
    font_size=None,
    background_color=None,
    horizontal_alignment=None,
    number_format=None
):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    for row in sheet[cell_range]:
        for cell in row:

            if bold:
                cell.font = Font(
                    bold=True,
                    size=font_size if font_size else 11
                )
            elif font_size:
                cell.font = Font(size=font_size)

            if background_color:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=background_color
                )

            if horizontal_alignment:
                cell.alignment = Alignment(
                    horizontal=horizontal_alignment
                )

            if number_format:
                cell.number_format = number_format

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "sheet": sheet_name,
        "range": cell_range
    }


def add_borders(file_path, sheet_name, cell_range):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in sheet[cell_range]:
        for cell in row:
            cell.border = border

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "sheet": sheet_name,
        "range": cell_range
    }


def auto_width(file_path, sheet_name):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                length = len(str(cell.value))
                if length > max_length:
                    max_length = length

        sheet.column_dimensions[column_letter].width = max_length + 2

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "sheet": sheet_name,
        "message": "Column widths adjusted"
    }


def sort_data(file_path, sheet_name, column_number, descending=False):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    rows.sort(
        key=lambda row: (
            row[column_number - 1]
            if row[column_number - 1] is not None
            else ""
        ),
        reverse=descending
    )

    for row_number, row_data in enumerate(rows, start=2):
        for column, value in enumerate(row_data, start=1):
            sheet.cell(row=row_number, column=column, value=value)

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "sheet": sheet_name,
        "column": column_number,
        "descending": descending
    }


def create_bar_chart(file_path, sheet_name, data_start_cell, data_end_cell, title):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    start_row = sheet[data_start_cell].row
    start_column = sheet[data_start_cell].column
    end_row = sheet[data_end_cell].row
    end_column = sheet[data_end_cell].column

    if end_column - start_column < 1:
        raise ValueError("Bar chart requires at least two columns.")

    chart = BarChart()
    data = Reference(sheet, min_col=start_column + 1, min_row=start_row,
                     max_col=end_column, max_row=end_row)
    categories = Reference(sheet, min_col=start_column,
                           min_row=start_row + 1, max_row=end_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.title = title

    sheet.add_chart(chart, "H2")
    _cloud_save(workbook, file_path)

    return {"success": True, "chart": "bar", "title": title}


def create_pie_chart(file_path, sheet_name, data_end_cell, title):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    end_row = sheet[data_end_cell].row
    end_column = sheet[data_end_cell].column

    if end_column < 2:
        raise ValueError("Pie chart requires at least two columns.")

    chart = PieChart()
    data = Reference(sheet, min_col=2, min_row=1, max_col=end_column, max_row=end_row)
    labels = Reference(sheet, min_col=1, min_row=2, max_row=end_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    chart.title = title

    sheet.add_chart(chart, "H2")
    _cloud_save(workbook, file_path)

    return {"success": True, "chart": "pie", "title": title}


def create_line_chart(file_path, sheet_name, data_end_cell, title):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    end_row = sheet[data_end_cell].row
    end_column = sheet[data_end_cell].column

    if end_column < 2:
        raise ValueError("Line chart requires at least two columns.")

    chart = LineChart()
    data = Reference(sheet, min_col=2, min_row=1, max_col=end_column, max_row=end_row)
    categories = Reference(sheet, min_col=1, min_row=2, max_row=end_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.title = title

    sheet.add_chart(chart, "H20")
    _cloud_save(workbook, file_path)

    return {"success": True, "chart": "line", "title": title}


def create_table(file_path, sheet_name, table_range, table_name="AITable"):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    start_cell, end_cell = table_range.split(":")

    min_row = sheet[start_cell].row
    max_row = sheet[end_cell].row
    min_col = sheet[start_cell].column
    max_col = sheet[end_cell].column

    if max_row < min_row or max_col < min_col:
        raise ValueError(f"Invalid table range: {table_range}")

    headers = []
    for column in range(min_col, max_col + 1):
        value = sheet.cell(row=min_row, column=column).value
        if (value is None or not isinstance(value, str) or not value.strip()):
            raise ValueError(
                f"Invalid table header at column {column}. "
                "Every table column must have a non-empty string header."
            )
        headers.append(value)

    existing_table_names = set()
    for existing_sheet in workbook.worksheets:
        for existing_table in existing_sheet.tables.values():
            existing_table_names.add(existing_table.name)

    if table_name in existing_table_names:
        raise ValueError(f"Table name already exists: {table_name}")

    table = Table(displayName=table_name, ref=table_range)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )
    table.tableStyleInfo = style

    sheet.add_table(table)
    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "table": table_name,
        "range": table_range,
        "headers": headers
    }


def clear_range(file_path, sheet_name, cell_range):
    workbook = load_workbook(file_path)
    sheet = workbook[sheet_name]

    for row in sheet[cell_range]:
        for cell in row:
            cell.value = None

    _cloud_save(workbook, file_path)

    return {
        "success": True,
        "sheet": sheet_name,
        "range": cell_range,
        "message": "Range cleared"
    }