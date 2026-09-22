# main.py
from typing import Optional
import asyncio
import json
import os
import time

from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware

from session_store import (
    create_session, get_session, add_message, add_file,
    set_document_context, set_agent_result, save_to_folder,
    update_session,
)
from file_processor import save_uploaded_file
from universal_file_reader import read_any_file
from dynamic_pdf_analyzer import analyze_any_pdf
from document_classifier import classify_document
from extractors.invoice import extract_invoice
from quick_calc import parse_calc
from brain import ask_ai, brain_status
from memory_helper import build_chat_prompt
from permission_manager import (
    create_permission_request, get_permission_request,
    approve_permission, reject_permission,
)
from planner import build_plan_for_task
from executor import execute_approved_plan
from agent import set_workbook_path
from rag_memory import add_document, build_context, get_all_documents

from workbook_store import safe_save
from excel_tool import set_cloud_context


app = FastAPI(title="Orbit Backend", version="10.1")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/")
def root():
    return {"status": "online", "service": "Orbit", "version": "10.1"}


@app.get("/brain/status")
def get_brain_status():
    return brain_status()


# ============================================================
# WORKBOOK READER
# ============================================================

def read_workbook(session, preferred_sheet=None):
    path = session.get("workbook_path")
    if not path or not os.path.exists(path):
        return {
            "id": int(time.time() * 1000),
            "name": "Workbook.xlsx",
            "columns": ["A", "B", "C", "D"],
            "rows": [["", "", "", ""]],
            "sheets": ["Sheet1"],
            "active_sheet": "Sheet1",
        }
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(path)
        active_from_session = session.get("active_sheet")

        if preferred_sheet and preferred_sheet in wb.sheetnames:
            ws = wb[preferred_sheet]
        elif active_from_session and active_from_session in wb.sheetnames:
            ws = wb[active_from_session]
        elif wb.sheetnames:
            ws = wb[wb.sheetnames[0]]
        else:
            ws = wb.active

        rows = []
        max_col = 0
        for row in ws.iter_rows(values_only=True):
            cleaned = ["" if v is None else v for v in row]
            if any(str(v).strip() for v in cleaned):
                max_col = max(max_col, len(cleaned))
                rows.append(cleaned)

        for row in rows:
            while len(row) < max_col:
                row.append("")

        columns = [get_column_letter(i + 1) for i in range(max_col)]
        if max_col == 0:
            columns = ["A", "B", "C", "D"]

        if rows:
            first = rows[0]
            if all(str(v).strip() for v in first):
                columns = [str(v) for v in first]
                rows = rows[1:]

        return {
            "id": int(time.time() * 1000),
            "name": session.get("workbook_name", "Workbook.xlsx"),
            "templateId": "session",
            "columns": columns,
            "rows": rows,
            "sheets": wb.sheetnames,
            "active_sheet": ws.title,
        }
    except Exception as e:
        print("read_workbook error:", e)
        return {
            "id": int(time.time() * 1000),
            "name": "Workbook.xlsx",
            "columns": ["A", "B", "C", "D"],
            "rows": [["", "", "", ""]],
            "sheets": ["Sheet1"],
            "active_sheet": "Sheet1",
        }


# ============================================================
# HELPERS
# ============================================================

def is_greeting(text):
    return text.lower().strip() in {
        "hi", "hello", "hey", "hii", "helo", "yo",
        "good morning", "good afternoon", "good evening", "good night",
    }


def is_work_request(text):
    text = (text or "").lower()
    kws = [
        "pdf", "excel", "sheet", "worksheet", "workbook", "cell",
        "row", "column", "table", "data", "create", "add", "delete",
        "update", "edit", "insert", "fill", "put", "enter", "extract",
        "calculate", "formula", "total", "average", "sum", "count",
        "sort", "filter", "format", "save", "export", "details", "work",
        "student", "employee", "customer", "invoice", "report", "bold",
        "mean", "max", "min", "percentage",
    ]
    return any(k in text for k in kws)


def is_plain_chat(text):
    return is_greeting(text) or not is_work_request(text)


def build_agent_task(message, session):
    analysis = session.get("last_document_analysis") or {}
    columns = analysis.get("columns") or []
    rows = analysis.get("rows") or []
    sheet_name = analysis.get("sheet_name") or "ExtractedData"

    rag_context = ""
    try:
        rag_context = build_context(session["session_id"], message, top_k=3)
    except Exception as e:
        print("[main] RAG error:", e)

    if columns and rows:
        clean_rows = []
        for r in rows:
            if isinstance(r, list):
                cleaned = [str(v).strip() if v is not None else "" for v in r]
                if len(cleaned) > 0:
                    clean_rows.append(cleaned)

        if clean_rows:
            task = (
                f'Create a new sheet called "{sheet_name}" '
                f'with columns {json.dumps(columns)} '
                f'and add these rows: {json.dumps(clean_rows)}.'
            )
            if rag_context:
                task += "\n\n" + rag_context
            return task

    conv = "\n".join(
        f"{m['role'].upper()}: {m['content']}"
        for m in session.get("messages", [])[-8:]
    ) or "No previous conversation."

    task = f"{message}\n\nConversation:\n{conv}"
    if rag_context:
        task += "\n\n" + rag_context
    return task


# ============================================================
# SESSION ENDPOINTS
# ============================================================

@app.post("/excel/session")
def create_excel_session(workbook: dict):
    try:
        sid = create_session(workbook)
        session = get_session(sid)
        default_sheet = (workbook.get("sheets") or ["Sheet1"])[0]
        update_session(sid, {"active_sheet": default_sheet})
        return {"success": True, "session_id": sid, "workbook": read_workbook(session)}
    except Exception as e:
        return {"success": False, "error": str(e)}


@app.get("/excel/session/{session_id}")
def read_excel_session(session_id: str):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}
    return {
        "session_id": session_id,
        "messages": session.get("messages", []),
        "files": session.get("files", []),
        "workbook": read_workbook(session),
        "document_analysis": session.get("last_document_analysis"),
        "agent_result": session.get("last_agent_result"),
        "company_folders": session.get("company_folders", {}),
        "active_sheet": session.get("active_sheet", "Sheet1"),
    }


# ============================================================
# SET ACTIVE SHEET
# ============================================================

@app.post("/excel/session/{session_id}/set-active-sheet")
def set_active_sheet_endpoint(session_id: str, sheet_name: str = Form(...)):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}
    update_session(session_id, {"active_sheet": sheet_name})
    print(f"[main] ✅ Active sheet set to: {sheet_name}")
    return {"success": True, "active_sheet": sheet_name}


# ============================================================
# READ SINGLE SHEET
# ============================================================

@app.get("/excel/session/{session_id}/sheet/{sheet_name}")
def read_single_sheet(session_id: str, sheet_name: str):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    path = session.get("workbook_path")
    if not path or not os.path.exists(path):
        return {"success": False, "error": "Workbook not found"}

    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(path)

        if sheet_name not in wb.sheetnames:
            return {"success": False, "error": f"Sheet not found: {sheet_name}"}

        ws = wb[sheet_name]

        rows = []
        max_col = 0
        for row in ws.iter_rows(values_only=True):
            cleaned = ["" if v is None else v for v in row]
            if any(str(v).strip() for v in cleaned):
                max_col = max(max_col, len(cleaned))
                rows.append(cleaned)

        for row in rows:
            while len(row) < max_col:
                row.append("")

        columns = [get_column_letter(i + 1) for i in range(max_col)]
        if max_col == 0:
            columns = ["A", "B", "C", "D"]

        if rows:
            first = rows[0]
            if all(str(v).strip() for v in first):
                columns = [str(v) for v in first]
                rows = rows[1:]

        return {
            "success": True,
            "sheet_name": sheet_name,
            "columns": columns,
            "rows": rows,
            "sheets": wb.sheetnames,
        }

    except Exception as e:
        print("read_single_sheet error:", e)
        return {"success": False, "error": str(e)}


# ============================================================
# ANALYSIS
# ============================================================

@app.get("/excel/session/{session_id}/analysis")
def analyze_workbook(session_id: str, sheet_name: Optional[str] = None):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    path = session.get("workbook_path")
    if not path or not os.path.exists(path):
        return {"success": False, "error": "Workbook not found"}

    active = sheet_name or session.get("active_sheet", "Sheet1")

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        if active not in wb.sheetnames:
            active = wb.sheetnames[0] if wb.sheetnames else "Sheet1"
        ws = wb[active]

        columns = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=1, column=c).value
            if v:
                columns.append(str(v))

        rows = []
        for r in range(2, ws.max_row + 1):
            row = []
            for c in range(1, len(columns) + 1):
                row.append(ws.cell(row=r, column=c).value)
            if any(v is not None for v in row):
                rows.append(row)

        analysis = {
            "sheet_name": active,
            "row_count": len(rows),
            "column_count": len(columns),
            "columns": columns,
            "numeric_columns": {},
            "text_columns": [],
            "warnings": [],
        }

        for col_idx, col_name in enumerate(columns):
            values = [row[col_idx] for row in rows if row[col_idx] is not None]
            if not values:
                continue

            numeric = [v for v in values if isinstance(v, (int, float))]
            if numeric and len(numeric) >= len(values) * 0.7:
                analysis["numeric_columns"][col_name] = {
                    "total": sum(numeric),
                    "average": sum(numeric) / len(numeric),
                    "min": min(numeric),
                    "max": max(numeric),
                    "count": len(numeric),
                }
            else:
                counts = {}
                for v in values:
                    if v is not None:
                        key = str(v)[:30]
                        counts[key] = counts.get(key, 0) + 1
                top = sorted(counts.items(), key=lambda x: -x[1])[:5]
                analysis["text_columns"].append({"name": col_name, "top_values": top})

        return {"success": True, "analysis": analysis}

    except Exception as e:
        print("analysis error:", e)
        return {"success": False, "error": str(e)}


# ============================================================
# QUICK CALC
# ============================================================

@app.post("/excel/session/{session_id}/quick-calc")
async def quick_calc_endpoint(session_id: str, message: str = Form(...)):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    path = session.get("workbook_path")
    if not path or not os.path.exists(path):
        return {"success": False, "error": "Workbook not found"}

    active = session.get("active_sheet", "Sheet1")

    set_cloud_context(session_id, session.get("workbook_name"))

    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(path)
        if active not in wb.sheetnames:
            active = wb.sheetnames[0] if wb.sheetnames else "Sheet1"
        ws = wb[active]

        columns = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=1, column=c).value
            if v:
                columns.append(str(v))

        if not columns:
            return {"success": False, "fallback": True}

        parsed = parse_calc(message, columns)
        if not parsed.get("matched"):
            return {"success": False, "fallback": True}

        if parsed.get("need_column"):
            return {
                "success": False,
                "need_column": True,
                "op": parsed["op"],
                "columns": columns,
                "message": f"Which column should I {parsed['op']}?",
            }

        op = parsed["op"]
        col_name = parsed["column"]

        col_idx = None
        for i, c in enumerate(columns, start=1):
            if str(c) == str(col_name):
                col_idx = i
                break
        if col_idx is None:
            return {"success": False, "fallback": True}

        last_row = ws.max_row
        letter = get_column_letter(col_idx)
        range_str = f"{letter}2:{letter}{last_row}"

        values = []
        for r in range(2, last_row + 1):
            v = ws.cell(row=r, column=col_idx).value
            if isinstance(v, (int, float)):
                values.append(v)

        if not values:
            return {"success": False, "message": f"No numbers found in column '{col_name}'"}

        fn_map = {"sum": "SUM", "average": "AVERAGE", "count": "COUNT", "max": "MAX", "min": "MIN"}
        formula = f"={fn_map[op]}({range_str})"

        if op == "sum": result = sum(values)
        elif op == "average": result = sum(values) / len(values)
        elif op == "count": result = len(values)
        elif op == "max": result = max(values)
        elif op == "min": result = min(values)
        else: return {"success": False, "fallback": True}

        target_row = last_row + 1
        label_col = col_idx - 1 if col_idx > 1 else col_idx
        if not ws.cell(row=target_row, column=label_col).value:
            ws.cell(row=target_row, column=label_col).value = op.capitalize()
        ws.cell(row=target_row, column=col_idx).value = formula

        safe_save(wb, path, session_id=session_id, filename=session.get("workbook_name"))

        print(f"[main] quick_calc op={op} col={col_name} range={range_str} value={result}")

        return {
            "success": True,
            "message": f"✅ {op.capitalize()} of {col_name} = {result:,.2f} → written below",
            "workbook": read_workbook(session, preferred_sheet=active),
            "op": op,
            "column": col_name,
            "value": result,
        }

    except Exception as e:
        print("quick_calc error:", e)
        return {"success": False, "error": str(e)}


# ============================================================
# DOWNLOAD WORKBOOK
# ============================================================

@app.get("/excel/session/{session_id}/download")
def download_workbook(session_id: str):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    path = session.get("workbook_path")
    filename = session.get("workbook_name", "Orbit.xlsx")

    if not path or not os.path.exists(path):
        try:
            from file_store import download_file
            remote_path = f"{session_id}/{filename}"
            local_tmp = os.path.join("sessions", session_id, filename)
            dl = download_file(remote_path, local_tmp)
            if dl.get("success"):
                path = local_tmp
            else:
                return {"success": False, "error": "Workbook not found in local or cloud"}
        except Exception as e:
            return {"success": False, "error": f"Cloud fetch failed: {e}"}

    if not filename.endswith(".xlsx"):
        filename += ".xlsx"

    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
    )


# ============================================================
# SAVE TO COMPANY FOLDER
# ============================================================

@app.post("/excel/session/{session_id}/save-to-folder")
def save_to_company_folder(session_id: str, folder: str = Form("Data"), filename: str = Form("")):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}
    path = session.get("workbook_path")
    if not path or not os.path.exists(path):
        return {"success": False, "error": "Workbook not found"}
    if not filename:
        filename = session.get("workbook_name", "Orbit.xlsx")
        if not filename.endswith(".xlsx"):
            filename += ".xlsx"
    return save_to_folder(session_id, folder, filename, path)


# ============================================================
# PLAN ENDPOINT
# ============================================================

@app.post("/excel/session/{session_id}/plan")
async def plan_task(session_id: str, message: str = Form(""), file: Optional[UploadFile] = File(None)):
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    message = (message or "").strip()

    print("\n" + "#" * 70)
    print("PLAN REQUEST  session:", session_id)
    print("MESSAGE:", message)
    print("FILE:", file.filename if file else "None")
    print("ACTIVE SHEET:", session.get("active_sheet", "Sheet1"))
    print("WORKBOOK:", session.get("workbook_path"))
    print("#" * 70)

    if message:
        add_message(session_id, "user", message)

    set_cloud_context(session_id, session.get("workbook_name"))

    if file:
        try:
            contents = await file.read()
            if not contents:
                return {"success": False, "status": "error", "message": "Empty file."}

            from io import BytesIO
            up = UploadFile(filename=file.filename, file=BytesIO(contents))
            uploaded = await save_uploaded_file(up)
            add_file(session_id, uploaded["filename"], uploaded)

            read_result = read_any_file(uploaded["path"])
            if not read_result.get("success") or not read_result.get("text"):
                return {"success": False, "status": "error", "message": f"Could not read file: {read_result.get('error')}"}

            document = {
                "has_text": True,
                "page_count": read_result.get("page_count", 1),
                "total_characters": len(read_result["text"]),
                "combined_text": read_result["text"],
                "file_type": read_result.get("type"),
            }

            classification = classify_document(read_result["text"], filename=uploaded["filename"])
            doc_type = classification.get("type", "generic")

            print(f"[main] Classified '{uploaded['filename']}' as: {doc_type} (confidence={classification.get('confidence')})")

            if doc_type == "invoice":
                analysis = extract_invoice(read_result["text"], filename=uploaded["filename"])
                print(f"[main] Invoice extractor used. Success={analysis.get('success')}")
            else:
                analysis = analyze_any_pdf(read_result["text"], filename=uploaded["filename"])

            if not analysis.get("success"):
                return {"success": False, "status": "error", "message": f"Could not analyze file: {analysis.get('error')}"}

            set_document_context(session_id, document, analysis, uploaded)
            add_document(session_id, uploaded["filename"], read_result["text"], analysis)

            if not message:
                all_docs = get_all_documents(session_id)
                return {
                    "success": True,
                    "status": "completed",
                    "action": "document",
                    "message": (
                        f"I analyzed '{uploaded['filename']}'.\n"
                        f"Type: {analysis.get('document_type')}\n"
                        f"Columns: {len(analysis.get('columns', []))}\n"
                        f"Rows: {len(analysis.get('rows', []))}\n"
                        f"Files in memory: {len(all_docs)}\n\n"
                        "Tell me what to do with it."
                    ),
                    "document_analysis": analysis,
                    "all_documents": all_docs,
                }
        except Exception as e:
            return {"success": False, "status": "error", "message": str(e)}

    if is_plain_chat(message):
        reply = ask_ai(build_chat_prompt(message, session))
        add_message(session_id, "assistant", reply)
        return {"success": True, "action": "conversation", "status": "completed", "message": reply}

    print("\n--- BUILDING PLAN ---")
    set_workbook_path(session["workbook_path"])

    built = await asyncio.to_thread(build_plan_for_task, message, session["workbook_path"], session)

    if not built.get("success"):
        return {"success": False, "status": "error", "message": built.get("message", "Could not build a plan.")}

    if built.get("already_satisfied"):
        return {"success": True, "already_satisfied": True, "summary": built.get("summary", "Already done.")}

    pid = create_permission_request(session_id, built["plan"], built["summary"], message)

    return {
        "success": True,
        "status": "awaiting_permission",
        "action": "excel",
        "permission_required": True,
        "permission_id": pid,
        "summary": built["summary"],
        "plan": built["plan"],
        "message": "Orbit prepared a plan. Please approve.",
    }


# ============================================================
# APPROVE
# ============================================================

@app.post("/excel/permission/{permission_id}/approve")
async def approve_plan(permission_id: str):
    req = get_permission_request(permission_id)
    if not req:
        return {"success": False, "error": "Permission not found"}
    if req["status"] != "pending":
        return {"success": False, "error": f"Already {req['status']}"}

    approve_permission(permission_id)
    session_id = req["session_id"]
    session = get_session(session_id)
    if not session:
        return {"success": False, "error": "Session not found"}

    plan = req["plan"]

    print("\n===== APPROVED PLAN =====")
    print(json.dumps(plan, indent=2, ensure_ascii=False))
    print("WORKBOOK:", session["workbook_path"])

    set_cloud_context(session_id, session.get("workbook_name"))

    exec_result = await asyncio.to_thread(
        execute_approved_plan, plan, session["workbook_path"], session_id, session.get("workbook_name"),
    )

    preferred_sheet = None
    for step in plan.get("steps", []):
        s = (step.get("arguments") or {}).get("sheet_name")
        if s:
            preferred_sheet = s
            break

    if preferred_sheet:
        update_session(session_id, {"active_sheet": preferred_sheet})

    session = get_session(session_id)
    workbook = read_workbook(session, preferred_sheet=preferred_sheet)
    set_agent_result(session_id, exec_result)

    steps = [{"id": "received", "label": "Task received", "status": "completed"}]
    for i, s in enumerate(exec_result.get("executed_steps", [])):
        ok = (s.get("result") or {}).get("success", False)
        steps.append({
            "id": f"step-{i}",
            "label": s.get("tool", "step").replace("_", " ").title(),
            "status": "completed" if ok else "error",
        })

    if exec_result.get("success"):
        steps.append({"id": "verified", "label": "Verified", "status": "completed"})
        message = "Orbit executed the approved plan."
    else:
        steps.append({"id": "failed", "label": "Failed", "status": "error"})
        message = exec_result.get("error", "Execution failed.")

    add_message(session_id, "assistant", message)

    return {
        "success": exec_result.get("success", False),
        "status": "completed" if exec_result.get("success") else "error",
        "message": message,
        "workbook": workbook,
        "steps": steps,
        "executed_steps": exec_result.get("executed_steps", []),
        "active_sheet": preferred_sheet or session.get("active_sheet"),
    }


# ============================================================
# REJECT
# ============================================================

@app.post("/excel/permission/{permission_id}/reject")
def reject_plan(permission_id: str):
    req = get_permission_request(permission_id)
    if not req:
        return {"success": False, "error": "Permission not found"}
    reject_permission(permission_id)
    return {"success": True, "status": "rejected", "message": "You rejected the plan. Nothing was changed."}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)