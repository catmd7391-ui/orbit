# universal_file_reader.py
import os


def detect_file_type(path):
    if not path:
        return "unknown"
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return "pdf"
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp"):
        return "image"
    if ext in (".xlsx", ".xls", ".csv"):
        return "excel"
    if ext in (".docx", ".doc"):
        return "word"
    if ext in (".txt", ".md", ".json", ".xml", ".html", ".htm"):
        return "text"
    return "unknown"


def read_any_file(path):
    if not path or not os.path.exists(path):
        return {"success": False, "error": "File not found", "text": "", "type": "unknown"}
    ftype = detect_file_type(path)
    try:
        if ftype == "pdf":
            return _read_pdf(path)
        if ftype == "text":
            return _read_text(path)
        if ftype == "excel":
            return _read_excel(path)
        if ftype == "word":
            return _read_word(path)
        if ftype == "image":
            return _read_image(path)
        return {"success": False, "error": f"Unsupported: {ftype}", "text": "", "type": ftype}
    except Exception as e:
        return {"success": False, "error": str(e), "text": "", "type": ftype}


def _read_pdf(path):
    text_parts = []
    page_count = 0
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            for i, page in enumerate(pdf.pages, start=1):
                t = page.extract_text() or ""
                if t.strip():
                    text_parts.append(f"--- PAGE {i} ---\n{t}")
    except Exception as e:
        print("pdfplumber failed:", e)

    full = "\n\n".join(text_parts).strip()

    if not full:
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            page_count = len(reader.pages)
            parts = []
            for i, page in enumerate(reader.pages, start=1):
                t = page.extract_text() or ""
                if t.strip():
                    parts.append(f"--- PAGE {i} ---\n{t}")
            full = "\n\n".join(parts).strip()
        except Exception as e:
            print("pypdf failed:", e)

    return {
        "success": bool(full),
        "type": "pdf",
        "text": full,
        "page_count": page_count,
        "error": None if full else "No readable text in PDF.",
    }


def _read_text(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()
    return {"success": bool(text.strip()), "type": "text", "text": text, "page_count": 1}


def _read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"--- SHEET: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    return {"success": True, "type": "excel", "text": text, "page_count": len(wb.worksheets)}


def _read_word(path):
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return {"success": bool(text.strip()), "type": "word", "text": text, "page_count": 1}
    except ImportError:
        return {"success": False, "error": "python-docx not installed", "text": "", "type": "word"}


def _read_image(path):
    try:
        import pytesseract
        from PIL import Image
        default = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.exists(default):
            pytesseract.pytesseract.tesseract_cmd = default
        img = Image.open(path)
        text = pytesseract.image_to_string(img)
        return {
            "success": bool(text.strip()),
            "type": "image",
            "text": text,
            "page_count": 1,
            "error": None if text.strip() else "No text in image",
        }
    except ImportError:
        return {"success": False, "error": "pytesseract/pillow not installed", "text": "", "type": "image"}